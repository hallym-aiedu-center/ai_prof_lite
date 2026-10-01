"""Deterministic local slide rendering and PPTX assembly helpers."""

from collections.abc import Sequence
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont, ImageOps
from pptx import Presentation
from pptx.util import Inches

PPT_IMAGE_WIDTH = 1920
PPT_IMAGE_HEIGHT = 1080
WIDTH = PPT_IMAGE_WIDTH
HEIGHT = PPT_IMAGE_HEIGHT

PPT_WIDTH_INCHES = 13.333333
PPT_HEIGHT_INCHES = 7.5

AVATAR_OVERLAY_HEIGHT = 520
AVATAR_OVERLAY_MARGIN = 34
AVATAR_SAFE_PADDING = 70
AVATAR_SAFE_MIN_WIDTH = 520
AVATAR_SAFE_MAX_WIDTH = 900


def _estimate_avatar_safe_zone(avatar_source_path: Path | None) -> dict:
    avatar_width = AVATAR_SAFE_MIN_WIDTH
    if avatar_source_path and avatar_source_path.is_file():
        try:
            with Image.open(avatar_source_path) as image:
                width, height = image.size
            if width > 0 and height > 0:
                avatar_width = round(AVATAR_OVERLAY_HEIGHT * (width / height))
        except (OSError, ValueError):
            avatar_width = AVATAR_SAFE_MIN_WIDTH

    avatar_width = max(360, min(avatar_width, AVATAR_SAFE_MAX_WIDTH))
    safe_width = min(
        PPT_IMAGE_WIDTH - 80,
        max(AVATAR_SAFE_MIN_WIDTH, avatar_width + AVATAR_SAFE_PADDING * 2),
    )
    safe_height = min(
        PPT_IMAGE_HEIGHT - 80, AVATAR_OVERLAY_HEIGHT + AVATAR_SAFE_PADDING * 2
    )
    right = PPT_IMAGE_WIDTH - AVATAR_OVERLAY_MARGIN
    bottom = PPT_IMAGE_HEIGHT - AVATAR_OVERLAY_MARGIN
    left = max(0, right - safe_width)
    top = max(0, bottom - safe_height)
    return {
        "left": left,
        "top": top,
        "right": right,
        "bottom": bottom,
        "width": right - left,
        "height": bottom - top,
        "width_pct": round((right - left) / PPT_IMAGE_WIDTH * 100),
        "height_pct": round((bottom - top) / PPT_IMAGE_HEIGHT * 100),
    }


def _contain_size(
    source_width: int, source_height: int, max_width: int, max_height: int
) -> tuple[int, int]:
    if min(source_width, source_height, max_width, max_height) <= 0:
        raise ValueError("Image dimensions must be positive.")
    scale = min(max_width / source_width, max_height / source_height)
    return max(1, round(source_width * scale)), max(1, round(source_height * scale))


def _fit_image(source_path: Path, box: tuple[int, int, int, int]) -> Image.Image:
    left, top, right, bottom = box
    max_width, max_height = right - left, bottom - top
    with Image.open(source_path) as source:
        image = ImageOps.exif_transpose(source).convert("RGB")
        width, height = _contain_size(image.width, image.height, max_width, max_height)
        return image.resize((width, height), Image.Resampling.LANCZOS)


def _font(size: int, *, bold: bool = False) -> ImageFont.ImageFont:
    candidates = [
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"
        if bold
        else "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/noto/NotoSansKR-Bold.ttf"
        if bold
        else "/usr/share/fonts/truetype/noto/NotoSansKR-Regular.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
        if bold
        else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).is_file():
            return ImageFont.truetype(candidate, size=size)
    return ImageFont.load_default()


def _wrap_text(
    draw: ImageDraw.ImageDraw, text: str, font: ImageFont.ImageFont, max_width: int
) -> list[str]:
    text = " ".join(str(text).split())
    if not text:
        return []
    units = text.split(" ")
    lines: list[str] = []
    current = ""
    for unit in units:
        candidate = unit if not current else f"{current} {unit}"
        if draw.textbbox((0, 0), candidate, font=font)[2] <= max_width:
            current = candidate
            continue
        if current:
            lines.append(current)
            current = ""
        fragment = ""
        for char in unit:
            candidate = fragment + char
            if fragment and draw.textbbox((0, 0), candidate, font=font)[2] > max_width:
                lines.append(fragment)
                fragment = char
            else:
                fragment = candidate
        current = fragment
    if current:
        lines.append(current)
    return lines


def _draw_slide_header(
    draw: ImageDraw.ImageDraw,
    *,
    lecture_title: str,
    slide_index: int,
    total_slides: int,
) -> None:
    draw.rectangle((0, 0, PPT_IMAGE_WIDTH, 18), fill=(79, 70, 229))
    draw.text(
        (90, 55), lecture_title[:100], font=_font(30, bold=True), fill=(100, 116, 139)
    )
    draw.text(
        (1690, 58),
        f"{slide_index:02d} / {total_slides:02d}",
        font=_font(26),
        fill=(100, 116, 139),
    )


def _draw_slide_text(
    draw: ImageDraw.ImageDraw,
    *,
    slide: dict,
    safe: dict,
    visual_exists: bool,
) -> None:
    title_font = _font(58, bold=True)
    body_font = _font(34)
    title = str(slide.get("title") or "핵심 내용").strip()
    y = 135
    for line in _wrap_text(draw, title, title_font, 1120)[:2]:
        draw.text((90, y), line, font=title_font, fill=(15, 23, 42))
        y += 76
    y += 35

    content_right = 1160 if visual_exists else min(1320, safe["left"] - 40)
    text_width = max(650, content_right - 155)
    bullets = [
        str(item).strip() for item in slide.get("bullets", []) if str(item).strip()
    ]
    for bullet in bullets[:6]:
        lines = _wrap_text(draw, bullet, body_font, text_width)
        if not lines:
            continue
        draw.ellipse((92, y + 14, 108, y + 30), fill=(79, 70, 229))
        for line in lines[:3]:
            draw.text((132, y), line, font=body_font, fill=(51, 65, 85))
            y += 49
        y += 20
        if y > 900:
            break


def _draw_supporting_visual(
    canvas: Image.Image,
    draw: ImageDraw.ImageDraw,
    *,
    visual_path: Path,
    safe: dict,
) -> None:
    box = (1220, 145, 1810, min(430, safe["top"] - 24))
    if box[2] - box[0] < 260 or box[3] - box[1] < 120:
        return
    visual = _fit_image(visual_path, box)
    x = box[0] + (box[2] - box[0] - visual.width) // 2
    y = box[1] + (box[3] - box[1] - visual.height) // 2
    draw.rounded_rectangle(
        (box[0] - 12, box[1] - 12, box[2] + 12, box[3] + 12),
        radius=28,
        fill=(255, 255, 255),
        outline=(226, 232, 240),
        width=3,
    )
    canvas.paste(visual, (x, y))


def _draw_presenter_safe_zone(draw: ImageDraw.ImageDraw, safe: dict) -> None:
    draw.rounded_rectangle(
        (safe["left"], safe["top"], safe["right"], safe["bottom"]),
        radius=40,
        fill=(241, 245, 249),
    )
    draw.text(
        (90, 1000), "AI Professor Lite", font=_font(24, bold=True), fill=(148, 163, 184)
    )


def render_local_slide(
    *,
    lecture_title: str,
    slide_index: int,
    total_slides: int,
    slide: dict,
    output_path: Path,
    visual_path: Path | None = None,
    avatar_source_path: Path | None = None,
) -> Path:
    canvas = Image.new("RGB", (PPT_IMAGE_WIDTH, PPT_IMAGE_HEIGHT), (248, 250, 252))
    draw = ImageDraw.Draw(canvas)
    safe = _estimate_avatar_safe_zone(avatar_source_path)
    visual_exists = visual_path is not None and visual_path.is_file()

    _draw_slide_header(
        draw,
        lecture_title=lecture_title,
        slide_index=slide_index,
        total_slides=total_slides,
    )
    _draw_slide_text(draw, slide=slide, safe=safe, visual_exists=visual_exists)
    if visual_exists:
        _draw_supporting_visual(canvas, draw, visual_path=visual_path, safe=safe)
    _draw_presenter_safe_zone(draw, safe)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(output_path, "PNG")
    return output_path


def build_pptx_from_images(
    *, output_path: Path, slide_image_paths: Sequence[Path]
) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(PPT_WIDTH_INCHES)
    prs.slide_height = Inches(PPT_HEIGHT_INCHES)
    for image_path in slide_image_paths:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(
            str(image_path), 0, 0, width=prs.slide_width, height=prs.slide_height
        )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)
    return output_path
