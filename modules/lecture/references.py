from __future__ import annotations

import asyncio
import math
import os
import re
import zipfile
from pathlib import Path
from uuid import uuid4
from defusedxml import ElementTree

from fastapi import HTTPException, UploadFile
from pptx import Presentation

from core.config import data_dir, positive_int
from core.openai.client import get_client
from core.openai.usage import embeddings_create

ALLOWED_REFERENCE_TYPES = {
    "application/pdf": ".pdf",
    "text/plain": ".txt",
    "text/markdown": ".md",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": ".pptx",
}
ALLOWED_SUFFIXES = {".pdf", ".txt", ".md", ".docx", ".pptx"}


def _safe_name(name: str | None, fallback_suffix: str) -> str:
    raw = Path(name or f"reference{fallback_suffix}").name
    suffix = fallback_suffix.lower()
    raw_stem = raw[:-len(suffix)] if suffix and raw.lower().endswith(suffix) else Path(raw).stem
    stem = re.sub(r"[^0-9A-Za-z가-힣._ -]+", "_", raw_stem).strip(" .") or "reference"
    maximum_stem = max(1, 160 - len(suffix))
    return f"{stem[:maximum_stem]}{suffix}"


async def save_reference_files(uploads: list[UploadFile] | None) -> list[dict]:
    uploads = [upload for upload in (uploads or []) if upload and upload.filename]
    if not uploads:
        return []

    maximum_files = positive_int("MAX_REFERENCE_FILES", 5)
    if len(uploads) > maximum_files:
        raise HTTPException(422, f"참고자료는 최대 {maximum_files}개까지 첨부할 수 있습니다.")

    maximum_total = positive_int("MAX_REFERENCE_UPLOAD_BYTES", 25 * 1024 * 1024)
    maximum_each = positive_int("MAX_REFERENCE_FILE_BYTES", 10 * 1024 * 1024)
    directory = data_dir() / "uploads" / "references" / uuid4().hex
    directory.mkdir(parents=True, exist_ok=False)
    saved: list[dict] = []
    total = 0

    try:
        for upload in uploads:
            content_type = (upload.content_type or "").split(";", 1)[0].strip().lower()
            suffix = Path(upload.filename or "").suffix.lower()
            expected_suffix = ALLOWED_REFERENCE_TYPES.get(content_type)
            if suffix not in ALLOWED_SUFFIXES or (expected_suffix and suffix != expected_suffix):
                raise HTTPException(422, "참고자료는 PDF, TXT, Markdown, DOCX, PPTX만 업로드할 수 있습니다.")

            content = bytearray()
            try:
                while chunk := await upload.read(64 * 1024):
                    content.extend(chunk)
                    if len(content) > maximum_each:
                        raise HTTPException(413, f"참고자료 1개는 최대 {maximum_each // (1024 * 1024)}MB입니다.")
                    if total + len(content) > maximum_total:
                        raise HTTPException(413, f"참고자료 전체 용량은 최대 {maximum_total // (1024 * 1024)}MB입니다.")
            finally:
                await upload.close()

            if not content:
                raise HTTPException(422, f"빈 참고자료는 사용할 수 없습니다: {upload.filename}")

            filename = _safe_name(upload.filename, suffix)
            path = directory / f"{uuid4().hex[:10]}_{filename}"
            temporary = path.with_suffix(path.suffix + ".tmp")
            try:
                await asyncio.to_thread(temporary.write_bytes, bytes(content))
                temporary.replace(path)
            finally:
                temporary.unlink(missing_ok=True)

            total += len(content)
            saved.append({
                "name": filename,
                "path": str(path),
                "content_type": content_type or "application/octet-stream",
                "size": len(content),
            })
        return saved
    except BaseException:
        for item in saved:
            Path(item["path"]).unlink(missing_ok=True)
        try:
            directory.rmdir()
        except OSError:
            pass
        raise


def cleanup_reference_files(files: list[dict] | None) -> None:
    parents: set[Path] = set()
    for item in files or []:
        path = Path(str(item.get("path") or ""))
        if path.is_file():
            path.unlink(missing_ok=True)
        if path.parent.name:
            parents.add(path.parent)
    for parent in parents:
        try:
            parent.rmdir()
        except OSError:
            pass


def _extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - installation guard
        raise RuntimeError("PDF 참고자료 처리를 위해 pypdf가 필요합니다.") from exc
    reader = PdfReader(str(path))
    return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages if (page.extract_text() or "").strip())


def _read_zip_member_limited(archive: zipfile.ZipFile, member: str, *, max_bytes: int) -> bytes:
    """Read one ZIP member without allowing unbounded decompression."""
    try:
        info = archive.getinfo(member)
    except KeyError as exc:
        raise ValueError(f"DOCX에 필수 파일이 없습니다: {member}") from exc

    if info.file_size > max_bytes:
        raise ValueError(
            f"DOCX 내부 XML 해제 크기가 허용 범위를 초과했습니다: "
            f"{info.file_size} > {max_bytes} bytes"
        )

    with archive.open(info, "r") as source:
        raw = source.read(max_bytes + 1)
    if len(raw) > max_bytes:
        raise ValueError(
            f"DOCX 내부 XML 해제 크기가 허용 범위를 초과했습니다: {max_bytes} bytes"
        )
    return raw


def _docx_paragraph_text(paragraph) -> str:
    parts: list[str] = []
    for node in paragraph.iter():
        local_name = node.tag.rsplit("}", 1)[-1]
        if local_name == "t" and node.text:
            parts.append(node.text)
        elif local_name == "tab":
            parts.append("\t")
        elif local_name in {"br", "cr"}:
            parts.append("\n")
    return "".join(parts).strip()


def _extract_docx(path: Path) -> str:
    maximum_xml_bytes = positive_int("MAX_DOCX_XML_BYTES", 8 * 1024 * 1024)
    with zipfile.ZipFile(path) as archive:
        raw = _read_zip_member_limited(
            archive,
            "word/document.xml",
            max_bytes=maximum_xml_bytes,
        )

    # defusedxml rejects entity expansion and other dangerous XML constructs.
    root = ElementTree.fromstring(raw)
    paragraphs = [
        text
        for paragraph in root.iter()
        if paragraph.tag.endswith("}p")
        for text in [_docx_paragraph_text(paragraph)]
        if text
    ]
    return "\n\n".join(paragraphs)


def _validate_pptx_archive_size(path: Path) -> None:
    maximum_xml_bytes = positive_int("MAX_PPTX_XML_BYTES", 8 * 1024 * 1024)
    maximum_uncompressed_bytes = positive_int(
        "MAX_PPTX_UNCOMPRESSED_BYTES",
        16 * 1024 * 1024,
    )

    with zipfile.ZipFile(path) as archive:
        infos = [info for info in archive.infolist() if not info.is_dir()]

        # Reject obvious zip bombs from the central directory before doing any
        # decompression. This covers media and embedded objects, not just XML.
        declared_total = 0
        declared_xml = 0
        for info in infos:
            declared_total += max(0, int(info.file_size))
            if declared_total > maximum_uncompressed_bytes:
                raise ValueError(
                    "PPTX 전체 해제 크기 합계가 허용 범위를 초과했습니다: "
                    f"{declared_total} > {maximum_uncompressed_bytes} bytes"
                )

            member = info.filename.lower()
            if member.endswith(".xml") or member.endswith(".rels"):
                declared_xml += max(0, int(info.file_size))
                if declared_xml > maximum_xml_bytes:
                    raise ValueError(
                        "PPTX 내부 XML 해제 크기 합계가 허용 범위를 초과했습니다: "
                        f"{declared_xml} > {maximum_xml_bytes} bytes"
                    )

        # Stream every member through a small buffer as a second bound. This
        # verifies the actual decompressed byte count without retaining media in
        # memory, before python-pptx is allowed to load the package.
        total_uncompressed = 0
        total_xml = 0
        for info in infos:
            member = info.filename.lower()
            is_xml = member.endswith(".xml") or member.endswith(".rels")
            with archive.open(info, "r") as source:
                while True:
                    remaining = maximum_uncompressed_bytes - total_uncompressed
                    chunk = source.read(min(64 * 1024, remaining + 1))
                    if not chunk:
                        break
                    total_uncompressed += len(chunk)
                    if total_uncompressed > maximum_uncompressed_bytes:
                        raise ValueError(
                            "PPTX 전체 해제 크기 합계가 허용 범위를 초과했습니다: "
                            f"> {maximum_uncompressed_bytes} bytes"
                        )
                    if is_xml:
                        total_xml += len(chunk)
                        if total_xml > maximum_xml_bytes:
                            raise ValueError(
                                "PPTX 내부 XML 해제 크기 합계가 허용 범위를 초과했습니다: "
                                f"> {maximum_xml_bytes} bytes"
                            )


def _extract_pptx(path: Path) -> str:
    _validate_pptx_archive_size(path)
    presentation = Presentation(path)
    chunks: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        texts = []
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if text and text.strip():
                texts.append(text.strip())
        if texts:
            chunks.append(f"[slide {index}]\n" + "\n".join(texts))
    return "\n\n".join(chunks)


def extract_reference_text(item: dict) -> str:
    path = Path(str(item.get("path") or ""))
    if not path.is_file():
        raise FileNotFoundError(f"참고자료 파일을 찾을 수 없습니다: {path}")
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix == ".pptx":
        return _extract_pptx(path)
    raise ValueError(f"지원하지 않는 참고자료 형식입니다: {suffix}")


def _chunks(text: str, *, size: int = 2600, overlap: int = 260) -> list[str]:
    """Split references on document boundaries before falling back to a sliding window.

    Embeddings work better when a chunk starts at a paragraph/slide boundary instead
    of an arbitrary Unicode character.  Oversized paragraphs still use overlap so
    context is not lost around a hard size boundary.
    """
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        return []
    if len(text) <= size:
        return [text]

    paragraphs = [part.strip() for part in re.split(r"\n\s*\n", text) if part.strip()]
    result: list[str] = []
    current = ""

    def flush_current() -> None:
        nonlocal current
        if current:
            result.append(current)
            current = ""

    for paragraph in paragraphs:
        if len(paragraph) <= size:
            candidate = paragraph if not current else f"{current}\n\n{paragraph}"
            if len(candidate) <= size:
                current = candidate
            else:
                flush_current()
                current = paragraph
            continue

        flush_current()
        step = max(1, size - overlap)
        for start in range(0, len(paragraph), step):
            piece = paragraph[start:start + size].strip()
            if piece:
                result.append(piece)
            if start + size >= len(paragraph):
                break

    flush_current()
    return result


def _cosine(a: list[float], b: list[float]) -> float:
    dot = sum(x * y for x, y in zip(a, b))
    na = math.sqrt(sum(x * x for x in a))
    nb = math.sqrt(sum(y * y for y in b))
    if not na or not nb:
        return 0.0
    return dot / (na * nb)


async def build_reference_context(
    files: list[dict] | None, *, title: str, topic: str, api_key: str,
    user_id: int, lecture_id: int | None = None,
) -> str:
    """Semantic RAG: chunk local files, embed query/chunks with OpenAI, rank by cosine similarity."""
    if not files:
        return ""

    documents: list[tuple[str, str]] = []
    for item in files:
        text = await asyncio.to_thread(extract_reference_text, item)
        text = text.strip()
        if text:
            documents.append((str(item.get("name") or Path(item["path"]).name), text))
    if not documents:
        return ""

    chunk_size = positive_int("REFERENCE_RAG_CHUNK_CHARS", 2600)
    try:
        overlap_value = int(os.getenv("REFERENCE_RAG_CHUNK_OVERLAP_CHARS", "260"))
    except ValueError as exc:
        raise ValueError("REFERENCE_RAG_CHUNK_OVERLAP_CHARS must be an integer") from exc
    overlap = min(max(0, overlap_value), max(0, chunk_size - 1))
    top_k = positive_int("REFERENCE_RAG_TOP_K", 16)
    try:
        max_chunks = int(os.getenv("REFERENCE_RAG_MAX_CHUNKS", "0"))
    except ValueError as exc:
        raise ValueError("REFERENCE_RAG_MAX_CHUNKS must be an integer") from exc
    if max_chunks < 0:
        raise ValueError(
            "REFERENCE_RAG_MAX_CHUNKS must be 0 (unlimited) or a positive integer"
        )
    limit = positive_int("MAX_REFERENCE_CONTEXT_CHARS", 60_000)
    model = os.getenv("REFERENCE_EMBEDDING_MODEL", "text-embedding-3-small").strip() or "text-embedding-3-small"

    chunks: list[tuple[str, str]] = []
    for name, text in documents:
        for chunk in _chunks(text, size=chunk_size, overlap=overlap):
            chunks.append((name, chunk))
            if max_chunks and len(chunks) >= max_chunks:
                break
        if max_chunks and len(chunks) >= max_chunks:
            break
    if not chunks:
        return ""

    query = f"강의 제목: {title}\n강의 주제: {topic}"
    client = get_client(api_key=api_key)
    vectors: list[list[float]] = []
    async with client:
        query_response = await embeddings_create(
            client, user_id=user_id, lecture_id=lecture_id,
            model=model, input=[query], encoding_format="float",
            usage_context={"operation": "reference_rag_embedding", "stage": "references.query", "item_key": "query"},
        )
        query_vector = list(query_response.data[0].embedding)
        batch_size = positive_int("REFERENCE_EMBEDDING_BATCH_SIZE", 64)
        for start in range(0, len(chunks), batch_size):
            batch = chunks[start:start + batch_size]
            response = await embeddings_create(
                client, user_id=user_id, lecture_id=lecture_id,
                model=model, input=[chunk for _, chunk in batch], encoding_format="float",
                usage_context={"operation": "reference_rag_embedding", "stage": "references.chunk_batch", "item_key": f"batch_{start // batch_size + 1}", "item_index": start // batch_size + 1, "metadata": {"batch_size": len(batch)}},
            )
            vectors.extend(list(item.embedding) for item in response.data)

    ranked = sorted(
        ((_cosine(query_vector, vector), index, chunks[index][0], chunks[index][1])
         for index, vector in enumerate(vectors)),
        key=lambda item: (-item[0], item[1]),
    )

    selected: list[str] = []
    used = 0
    for score, _, name, chunk in ranked[:top_k]:
        block = f"### {name} (semantic RAG, similarity={score:.4f})\n{chunk}"
        if used + len(block) > limit:
            remaining = limit - used
            if remaining > 200:
                selected.append(block[:remaining])
            break
        selected.append(block)
        used += len(block)
    return "\n\n".join(selected)
