from pathlib import Path
from unittest.mock import AsyncMock

from PIL import Image, ImageDraw
from pptx import Presentation

from modules.lecture import slides as slides_module
from modules.lecture.slides import (
    HEIGHT,
    WIDTH,
    _contain_size,
    _fit_image,
    build_slide_assets,
    render_local_slide,
)


def test_contain_size_preserves_generated_image_ratio():
    assert _contain_size(1536, 1024, 700, 635) == (700, 467)


def test_fit_image_contains_without_cropping(tmp_path: Path):
    source_path = tmp_path / "source.png"
    source = Image.new("RGB", (1536, 1024), "white")
    draw = ImageDraw.Draw(source)
    draw.rectangle((0, 0, 80, 1023), fill="red")
    draw.rectangle((1455, 0, 1535, 1023), fill="blue")
    source.save(source_path)

    fitted = _fit_image(source_path, (1110, 245, 1810, 880))
    assert fitted.size == (700, 467)
    assert fitted.getpixel((0, fitted.height // 2))[0] > 200
    assert fitted.getpixel((fitted.width - 1, fitted.height // 2))[2] > 200


async def test_pptx_uses_one_full_slide_image_per_page(tmp_path: Path, monkeypatch):
    monkeypatch.setenv("LECTURE_RENDER_SLIDES_WITH_IMAGE_MODEL", "false")
    visual_path = tmp_path / "visual.png"
    Image.new("RGB", (1200, 900), "white").save(visual_path)
    plan = {
        "slides": [
            {"title": "전체 슬라이드 이미지 테스트", "bullets": ["첫 번째 내용", "두 번째 내용"]},
            {"title": "두 번째 슬라이드", "bullets": ["이미지가 없어도 PNG 한 장입니다."]},
        ]
    }

    pptx_path, pngs = await build_slide_assets(
        api_key="unused",
        title="테스트 강의",
        plan=plan,
        output_dir=tmp_path / "slides",
        image_model="unused",
        generate_images=True,
        image_paths=[visual_path, None],
    )

    assert len(pngs) == 2
    for png in pngs:
        with Image.open(png) as rendered:
            assert rendered.size == (WIDTH, HEIGHT)

    prs = Presentation(pptx_path)
    assert len(prs.slides) == 2
    for slide in prs.slides:
        assert len(slide.shapes) == 1
        picture = slide.shapes[0]
        assert picture.shape_type == 13
        assert picture.left == 0 and picture.top == 0
        assert picture.width == prs.slide_width and picture.height == prs.slide_height
        assert picture.crop_left == picture.crop_right == picture.crop_top == picture.crop_bottom == 0


async def test_generate_images_false_never_calls_image_api(tmp_path: Path, monkeypatch):
    monkeypatch.setenv("LECTURE_RENDER_SLIDES_WITH_IMAGE_MODEL", "true")
    generated = AsyncMock()
    monkeypatch.setattr(slides_module, "generate_slide_image", generated)

    await build_slide_assets(
        api_key="must-not-be-used-for-images",
        title="No image API",
        plan={"slides": [{"title": "Local", "bullets": ["Only local rendering"]}]},
        output_dir=tmp_path / "slides",
        image_model="gpt-image-2",
        generate_images=False,
    )
    generated.assert_not_awaited()


def test_local_slide_renders_supporting_visual(tmp_path: Path):
    visual_path = tmp_path / "visual.png"
    Image.new("RGB", (800, 600), (240, 20, 20)).save(visual_path)
    output_path = tmp_path / "slide.png"

    render_local_slide(
        lecture_title="테스트 강의",
        slide_index=1,
        total_slides=1,
        slide={"title": "제목", "bullets": ["설명"]},
        output_path=output_path,
        visual_path=visual_path,
    )

    with Image.open(output_path) as rendered:
        # The supporting visual is placed above the lower-right presenter safe
        # zone. A solid red source makes the regression easy to detect.
        crop = rendered.crop((1220, 145, 1810, 362)).convert("RGB")
        red_pixels = sum(
            1 for r, g, b in crop.getdata()
            if r > 200 and g < 80 and b < 80
        )
    assert red_pixels > 1000
